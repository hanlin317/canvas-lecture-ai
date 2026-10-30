from __future__ import annotations

import os
import subprocess
import sys
import zipfile
from pathlib import Path
from subprocess import CompletedProcess
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document

from app import config as app_config
from app.ai import factory as ai_factory
from app.ai.base import AIRequest
from app.ai.factory import AI_PROVIDER_SPECS, api_key_status, get_ai_provider, load_api_key
from app.ai.openai_provider import OpenAIProvider
from app.ai.prompts import PROMPT_TEMPLATES, build_study_note_prompt, note_length_tokens
from app.canvas.client import CanvasClient, _next_page_url
from app.canvas.models import CanvasCourseSummary
from app.config import (
    application_support_dir,
    load_settings,
    prepare_runtime_environment,
    save_web_settings,
)
from app.database import Database
from app.downloads.manager import DownloadManager, ffmpeg_executable
from app.echo360.client import (
    Echo360Client,
    ResolvedEchoAssets,
    _best_track_url,
    _matching_transcript_url,
)
from app.models import (
    Asset,
    Course,
    DownloadStatus,
    EchoPresentationAsset,
    EchoRecording,
    JobItemRecord,
    JobRecord,
    Lecture,
    LectureContext,
    NoteDraft,
    PipelineOptions,
    ProcessingRecord,
    Topic,
)
from app.notebooklm.exporter import notebooklm_markdown
from app.obsidian.exporter import export_ai_note, obsidian_note_markdown, save_obsidian_note
from app.services.discovery import (
    DiscoveryService,
    _default_course_group,
    _normalize_course_code,
    _parse_course_groups,
)
from app.services.manifest import export_topic_manifest
from app.services.pipeline import PipelineService
from app.transcripts.cleaner import clean_plain_transcript, clean_vtt
from app.transcripts.whisper import transcribe_with_local_whisper, transcribe_with_mlx_whisper
from app.utils import lecture_stem, number_for_position, safe_filename, sha256_text


def test_safe_filename_for_macos() -> None:
    assert safe_filename("  3.1 Intro: Endo / canals \n ") == "3.1 Intro- Endo - canals"
    assert safe_filename("...") == "untitled"


def test_number_for_position() -> None:
    assert number_for_position("3", 1) == "3.1"
    assert number_for_position("3.", 2) == "3.2"
    assert number_for_position("", 4) == "4"


def test_lecture_stem_uses_course_topic_local_number_and_title() -> None:
    assert (
        lecture_stem("SDDM6321", "Periodontics", "1", "Periodontal Maintenance")
        == "SDDM6321 Periodontics 1 Periodontal Maintenance"
    )


def test_ffmpeg_executable_finds_packaged_windows_helper(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    helper = tmp_path / "ffmpeg" / "ffmpeg.exe"
    helper.parent.mkdir()
    helper.write_bytes(b"packaged ffmpeg")
    monkeypatch.setattr(sys, "_MEIPASS", str(tmp_path), raising=False)
    monkeypatch.setattr(sys, "platform", "win32")

    assert ffmpeg_executable() == str(helper)


def test_canvas_course_code_normalization() -> None:
    assert (
        _normalize_course_code("2026_SDDM6321", "SDDM6321 Core Clinical Practice 3") == "SDDM6321"
    )
    assert _normalize_course_code("2026_SDDM6311", "Core Clinical Practice 2") == "SDDM6311"
    assert (
        _normalize_course_code("2026_SDDM5418_SDDM5328", "DMD3 & DMD4 Electives")
        == "SDDM5418_SDDM5328"
    )
    assert _normalize_course_code("DATA1001", "Foundations of Data Science") == "DATA1001"


def test_course_group_fallbacks_and_ai_response_parsing() -> None:
    assert _default_course_group("SDDM6321", "Core Clinical Practice") == "Dentistry"
    assert _default_course_group("DATA1001", "Foundations") == "Data, Computing & Maths"
    response = """```json
    {"groups":[
      {"name":"Dentistry","canvas_ids":["1"]},
      {"name":"Data Science","canvas_ids":["2"]}
    ]}
    ```"""
    assert _parse_course_groups(response, {"1", "2"}) == {
        "1": "Dentistry",
        "2": "Data Science",
    }


def test_transcript_cleaning_vtt() -> None:
    raw = """WEBVTT

1
00:00:01.000 --> 00:00:03.000
Root canal

NOTE CONF {"raw":[99,100]}

2
00:00:02.000 --> 00:00:04.000
Root canal therapy.

STYLE
::cue { color: lime; }
"""
    assert clean_vtt(raw) == "Root canal therapy.\n"


def test_plain_transcript_cleaning() -> None:
    assert clean_plain_transcript(" hello   world \nhello   world\n") == "hello world\n"


def test_anonymized_transcript_fixture_removes_repeated_lines() -> None:
    fixture = Path(__file__).parent / "fixtures" / "2026-10-29-anonymized-transcript.txt"
    cleaned = clean_plain_transcript(fixture.read_text(encoding="utf-8"))

    assert cleaned.count("periodontal maintenance") == 1
    assert cleaned.count("The recall interval") == 1
    assert cleaned.count("The key point") == 1


def test_study_note_prompt_applies_length_and_custom_instructions() -> None:
    prompt = build_study_note_prompt(
        "SDDM6321",
        "Periodontics",
        "1",
        "Periodontal Maintenance",
        "Transcript evidence",
        note_length="detailed",
        custom_instructions="Emphasize maintenance intervals and clinical reasoning.",
    )

    assert "Detailed (approximately 3,000-4,500 words" in prompt
    assert "Emphasize maintenance intervals and clinical reasoning." in prompt
    assert "Transcript evidence" in prompt
    assert note_length_tokens("concise") == 2048
    assert note_length_tokens("comprehensive") == 12288


def test_study_note_prompt_supports_each_template() -> None:
    for template_id, template in PROMPT_TEMPLATES.items():
        prompt = build_study_note_prompt(
            "SDDM6321",
            "Periodontics",
            "1",
            "Maintenance",
            "Transcript evidence",
            prompt_template=template_id,
        )
        assert template["label"] in prompt
        assert template["guidance"] in prompt


@pytest.mark.asyncio
async def test_openai_provider_uses_direct_chat_completions_request(monkeypatch) -> None:
    captured: dict = {}

    class FakeResponse:
        def raise_for_status(self) -> None:
            return None

        def json(self) -> dict:
            return {"choices": [{"message": {"content": "Generated note"}}]}

    class FakeClient:
        def __init__(self, **_kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def post(self, url, *, headers, json):
            captured.update(url=url, headers=headers, json=json)
            return FakeResponse()

    monkeypatch.setattr("app.ai.openai_provider.httpx.AsyncClient", FakeClient)
    result = await OpenAIProvider("test-key").generate_note(
        AIRequest(
            model="gpt-test",
            system_prompt="System",
            user_prompt="User",
            max_output_tokens=1234,
        )
    )

    assert result == "Generated note"
    assert captured["url"] == "https://api.openai.com/v1/chat/completions"
    assert captured["headers"]["Authorization"] == "Bearer test-key"
    assert captured["json"]["max_completion_tokens"] == 1234


def test_database_persists_queue_profiles_and_note_drafts(tmp_path: Path) -> None:
    db = Database(tmp_path / "app.sqlite3")
    course = db.upsert_course(Course(canvas_id="1", code="TEST1000", name="Test"))
    lecture = db.upsert_lecture(
        Lecture(source_id="recording-1", course_id=course.id or 0, original_title="Lecture")
    )
    job = db.create_job(JobRecord(id="job-1", kind="pipeline", payload_json='{"ok":true}'))
    db.create_job_items(
        [JobItemRecord(job_id=job.id, lecture_id=lecture.id or 0, position=1)]
    )
    db.save_profile("course:1", {"privacy_mode": "local_only"}, "2026-10-29T00:00:00Z")
    processing = db.upsert_processing(
        ProcessingRecord(lecture_id=lecture.id or 0, stage="ai_note", status="complete")
    )
    draft = NoteDraft(
        processing_id=processing.id or 0,
        lecture_id=lecture.id or 0,
        markdown="# Editable",
    )
    db.upsert_note_draft(draft)

    assert db.get_job("job-1").payload_json == '{"ok":true}'
    assert db.list_job_items("job-1")[0].lecture_id == lecture.id
    assert db.get_profile("course:1") == {"privacy_mode": "local_only"}
    assert db.get_note_draft(processing.id or 0).markdown == "# Editable"


def test_database_reconciles_interrupted_pipeline_jobs(tmp_path: Path) -> None:
    db = Database(tmp_path / "app.sqlite3")
    course = db.upsert_course(Course(canvas_id="1", code="TEST1000", name="Test"))
    lecture = db.upsert_lecture(
        Lecture(source_id="recording-1", course_id=course.id or 0, original_title="Lecture")
    )
    db.create_job(JobRecord(id="job-restart", kind="pipeline", status="running"))
    db.create_job_items(
        [
            JobItemRecord(
                job_id="job-restart",
                lecture_id=lecture.id or 0,
                position=1,
                status="running",
            )
        ]
    )

    db.reconcile_interrupted_jobs()

    assert db.get_job("job-restart").status == "paused"
    assert db.list_job_items("job-restart")[0].status == "paused"


@pytest.mark.asyncio
async def test_download_resumes_an_existing_partial_file(tmp_path: Path, monkeypatch) -> None:
    destination = tmp_path / "lecture.mp4"
    destination.with_suffix(".mp4.part").write_bytes(b"first-")
    requested_headers: dict[str, str] = {}

    class FakeResponse:
        status_code = 206

        def raise_for_status(self) -> None:
            return None

        async def aiter_bytes(self, _size: int):
            yield b"second"

    class FakeStream:
        async def __aenter__(self):
            return FakeResponse()

        async def __aexit__(self, *_args):
            return None

    class FakeClient:
        def __init__(self, **_kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        def stream(self, _method, _url, *, headers):
            requested_headers.update(headers)
            return FakeStream()

    monkeypatch.setattr("app.downloads.manager.httpx.AsyncClient", FakeClient)

    result = await DownloadManager(retries=0).download("https://example.test/video", destination)

    assert result == "complete"
    assert requested_headers["Range"] == "bytes=6-"
    assert destination.read_bytes() == b"first-second"


def test_bundled_mlx_whisper_writes_local_transcript(tmp_path: Path, monkeypatch) -> None:
    video = tmp_path / "lecture.mp4"
    video.write_bytes(b"video")
    output = tmp_path / "lecture.whisper.txt"
    calls = []
    fake_module = SimpleNamespace(
        transcribe=lambda path, path_or_hf_repo: (
            calls.append((path, path_or_hf_repo)) or {"text": "Local transcript"}
        )
    )
    monkeypatch.setattr("app.transcripts.whisper.sys.platform", "darwin")
    monkeypatch.setattr("app.transcripts.whisper.importlib.util.find_spec", lambda _name: object())
    monkeypatch.setitem(sys.modules, "mlx_whisper", fake_module)

    result = transcribe_with_mlx_whisper(
        video,
        output,
        model="mlx-community/whisper-large-v3-turbo-q4",
    )

    assert result == output
    assert output.read_text(encoding="utf-8") == "Local transcript\n"
    assert calls == [(str(video), "mlx-community/whisper-large-v3-turbo-q4")]


def test_faster_whisper_writes_local_transcript_on_windows(tmp_path: Path, monkeypatch) -> None:
    video = tmp_path / "lecture.mp4"
    video.write_bytes(b"video")
    output = tmp_path / "lecture.whisper.txt"
    calls = []

    class FakeWhisperModel:
        def __init__(self, model, *, device, compute_type):
            calls.append((model, device, compute_type))

        def transcribe(self, path, *, vad_filter):
            calls.append((path, vad_filter))
            return iter([SimpleNamespace(text=" First "), SimpleNamespace(text="second")]), None

    monkeypatch.setattr("app.transcripts.whisper.sys.platform", "win32")
    monkeypatch.setattr(
        "app.transcripts.whisper.importlib.util.find_spec",
        lambda name: object() if name == "faster_whisper" else None,
    )
    monkeypatch.setitem(
        sys.modules,
        "ctranslate2",
        SimpleNamespace(get_cuda_device_count=lambda: 0),
    )
    monkeypatch.setitem(
        sys.modules,
        "faster_whisper",
        SimpleNamespace(WhisperModel=FakeWhisperModel),
    )

    result = transcribe_with_local_whisper(video, output, model="large-v3-turbo")

    assert result == output
    assert output.read_text(encoding="utf-8") == "First second\n"
    assert calls == [("large-v3-turbo", "cpu", "int8"), (str(video), True)]


def test_hash_is_content_based() -> None:
    assert sha256_text("abc") == sha256_text("abc")
    assert sha256_text("abc") != sha256_text("abcd")


def test_notebooklm_markdown() -> None:
    ctx = LectureContext("SDDM6321", "Endodontics", "3.1", "Intro", "Echo360", "abc")
    md = notebooklm_markdown(ctx, "Transcript body")
    assert "course: SDDM6321" in md
    assert "## Transcript" in md
    assert "Transcript body" in md


def test_obsidian_markdown() -> None:
    ctx = LectureContext("SDDM6321", "Endodontics", "3.1", "Intro", "Echo360", "abc")
    md = obsidian_note_markdown(
        ctx,
        "## Core concepts\nContent",
        parent_note="Endodontics",
        ai_provider="openai",
        ai_model="model",
        prompt_version="1.0",
        transcript="Transcript",
    )
    assert 'up: "[[Endodontics]]"' in md
    assert "transcript_hash:" in md
    assert "## Core concepts" in md


def test_obsidian_filename_uses_only_lecture_title(tmp_path: Path) -> None:
    ctx = LectureContext(
        "SDDM6321",
        "Endodontics",
        "3.1",
        "Access Cavities: Principles",
        "Echo360",
        "abc",
    )

    output = save_obsidian_note("# Note", tmp_path, ctx)

    assert output.name == "Access Cavities- Principles.md"


def test_config_parsing(tmp_path: Path) -> None:
    config = tmp_path / "config.toml"
    config.write_text(
        "[downloads]\nconcurrency = 2\n[canvas]\nbase_url = 'https://canvas.example'\n",
        encoding="utf-8",
    )
    settings = load_settings(config)
    assert settings.downloads.concurrency == 2
    assert settings.canvas.base_url == "https://canvas.example"


def test_release_notes_are_extracted_from_matching_dated_version(tmp_path: Path) -> None:
    changelog = tmp_path / "changelog.md"
    output = tmp_path / "notes.md"
    changelog.write_text(
        "# Changelog\n\n## [Unreleased]\n\n### Added\n\n- Future.\n\n"
        "## [1.2.0] - 2026-10-26\n\n### Added\n\n- Released feature.\n\n"
        "## [1.1.0] - 2026-10-20\n\n### Fixed\n\n- Earlier fix.\n",
        encoding="utf-8",
    )
    script = Path(__file__).parents[1] / "scripts" / "2026-10-26-release-notes.py"

    result = subprocess.run(
        [
            sys.executable,
            str(script),
            "--changelog",
            str(changelog),
            "--version",
            "1.2.0",
            "--repository",
            "owner/project",
            "--previous-version",
            "1.1.0",
            "--output",
            str(output),
        ],
        check=False,
        capture_output=True,
        text=True,
    )

    notes = output.read_text(encoding="utf-8")
    assert result.returncode == 0
    assert "Released feature" in notes
    assert "Earlier fix" not in notes
    assert "owner/project/compare/v1.1.0...v1.2.0" in notes
    assert "### Installation" in notes


def test_release_notes_reject_missing_version(tmp_path: Path) -> None:
    changelog = tmp_path / "changelog.md"
    output = tmp_path / "notes.md"
    changelog.write_text("# Changelog\n\n## [Unreleased]\n", encoding="utf-8")
    script = Path(__file__).parents[1] / "scripts" / "2026-10-26-release-notes.py"

    result = subprocess.run(
        [
            sys.executable,
            str(script),
            "--changelog",
            str(changelog),
            "--version",
            "9.9.9",
            "--repository",
            "owner/project",
            "--output",
            str(output),
        ],
        check=False,
        capture_output=True,
        text=True,
    )

    assert result.returncode != 0
    assert "does not contain a release entry for 9.9.9" in result.stderr
    assert not output.exists()


def test_web_settings_update_preserves_comments_and_unrelated_values(tmp_path: Path) -> None:
    config = tmp_path / "config.toml"
    config.write_text(
        "# Keep this comment\n"
        "[paths]\n"
        'download_root = "~/Documents" # downloads\n'
        'notebooklm_staging_dir = "~/Notebook"\n'
        'obsidian_vault = ""\n\n'
        "[downloads]\nconcurrency = 7\n",
        encoding="utf-8",
    )

    save_web_settings(
        {
            "app.onboarding_step": 4,
            "ai.enabled": False,
            "canvas.base_url": "https://canvas.example.edu",
            "paths.obsidian_vault": "/tmp/My Vault",
            "ai.provider": "anthropic",
            "ai.model": "claude-test",
            "transcripts.whisper_enabled": True,
            "transcripts.whisper_model": "whisper-test",
        },
        config,
    )

    text = config.read_text(encoding="utf-8")
    settings = load_settings(config)
    assert "# Keep this comment" in text
    assert "# downloads" in text
    assert settings.downloads.concurrency == 7
    assert settings.app.onboarding_step == 4
    assert settings.ai.enabled is False
    assert settings.canvas.base_url == "https://canvas.example.edu"
    assert settings.paths.obsidian_vault == "/tmp/My Vault"
    assert settings.ai.provider == "anthropic"
    assert settings.transcripts.whisper_enabled is True


def test_packaged_runtime_bootstraps_private_application_support(
    tmp_path: Path, monkeypatch
) -> None:
    bundle = tmp_path / "bundle"
    data_dir = tmp_path / "application-support"
    (bundle / "playwright-browsers").mkdir(parents=True)
    (bundle / "ffmpeg").mkdir()
    ffmpeg = bundle / "ffmpeg" / "ffmpeg-test"
    ffmpeg.write_text("binary", encoding="utf-8")
    (bundle / "config.example.toml").write_text(
        "[canvas]\nbase_url = 'https://canvas.example'\n",
        encoding="utf-8",
    )
    original_cwd = Path.cwd()
    monkeypatch.setenv("DLM_DATA_DIR", str(data_dir))
    monkeypatch.delenv("DLM_CONFIG", raising=False)
    monkeypatch.delenv("DLM_FFMPEG", raising=False)
    monkeypatch.delenv("PLAYWRIGHT_BROWSERS_PATH", raising=False)
    monkeypatch.setattr(app_config, "is_packaged", lambda: True)
    monkeypatch.setattr(app_config, "bundled_resource", lambda name: bundle / name)

    try:
        prepared = prepare_runtime_environment()
    finally:
        os.chdir(original_cwd)

    assert prepared == data_dir
    assert (data_dir / "config.toml").is_file()
    assert os.environ["DLM_CONFIG"] == str(data_dir / "config.toml")
    assert os.environ["PLAYWRIGHT_BROWSERS_PATH"] == str(bundle / "playwright-browsers")
    assert os.environ["DLM_FFMPEG"] == str(ffmpeg)


def test_windows_application_support_uses_local_app_data(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.delenv("DLM_DATA_DIR", raising=False)
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    monkeypatch.setattr(app_config.sys, "platform", "win32")

    assert application_support_dir() == tmp_path / "Canvas Lecture AI"


def test_canvas_recognizes_external_sso_as_expired() -> None:
    settings = load_settings()
    settings.canvas.base_url = "https://canvas.example.edu"
    client = CanvasClient(settings)

    assert client._is_authentication_url("https://sso.example.edu/app/canvas/sso/saml")
    assert client._is_authentication_url("https://canvas.example.edu/login/saml")
    assert not client._is_authentication_url("https://canvas.example.edu/courses/12345")
    assert client._is_authenticated_page(
        "https://canvas.example.edu/courses/12345",
        "Dashboard Courses Account",
    )


def test_canvas_pagination_link_parser() -> None:
    header = (
        '<https://canvas.example/api/v1/courses?page=1>; rel="current", '
        '<https://canvas.example/api/v1/courses?page=2>; rel="next", '
        '<https://canvas.example/api/v1/courses?page=4>; rel="last"'
    )

    assert _next_page_url(header) == "https://canvas.example/api/v1/courses?page=2"
    assert _next_page_url(None) is None


@pytest.mark.asyncio
async def test_canvas_token_api_follows_all_pages(monkeypatch) -> None:
    class FakeResponse:
        def __init__(self, payload, link="") -> None:
            self.payload = payload
            self.headers = {"link": link}

        def raise_for_status(self) -> None:
            return None

        def json(self):
            return self.payload

    calls = []
    responses = [
        FakeResponse(
            [{"id": 1}],
            '<https://canvas.example/api/v1/courses?page=2>; rel="next"',
        ),
        FakeResponse([{"id": 2}]),
    ]

    class FakeClient:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def get(self, url, params=None, headers=None):
            calls.append((url, params, headers))
            return responses.pop(0)

    settings = load_settings()
    settings.canvas.base_url = "https://canvas.example"
    monkeypatch.setenv("CANVAS_ACCESS_TOKEN", "token")
    monkeypatch.setattr("app.canvas.client.httpx.AsyncClient", lambda **_kwargs: FakeClient())

    records = await CanvasClient(settings)._api_get("/api/v1/courses", {"per_page": "100"})

    assert records == [{"id": 1}, {"id": 2}]
    assert calls[0][1] == {"per_page": "100"}
    assert calls[1][0].endswith("page=2")
    assert calls[1][1] is None


@pytest.mark.asyncio
async def test_course_sync_requests_every_accessible_course(tmp_path: Path) -> None:
    db = Database(tmp_path / "app.sqlite3")
    settings = load_settings()
    service = DiscoveryService(db, settings)
    service.canvas.list_courses = AsyncMock(
        return_value=[
            CanvasCourseSummary(
                canvas_id="1",
                code="2025_ARCH1000",
                name="Archived Course",
                url="https://canvas.example/courses/1",
            )
        ]
    )

    courses = await service.sync_canvas_courses()

    service.canvas.list_courses.assert_awaited_once_with(active_only=False)
    assert [course.code for course in courses] == ["2025_ARCH1000"]


def test_ai_provider_catalog_and_keychain_lookup(monkeypatch) -> None:
    names = {spec.name for spec in AI_PROVIDER_SPECS}
    assert {"openai", "anthropic", "gemini", "openrouter", "xai", "deepseek", "mistral"} <= names
    assert all(spec.models for spec in AI_PROVIDER_SPECS)
    assert all(spec.default_model in spec.model_ids for spec in AI_PROVIDER_SPECS)
    deepseek = next(spec for spec in AI_PROVIDER_SPECS if spec.name == "deepseek")
    assert deepseek.default_model == "deepseek-flash"
    assert any(item.label == "DeepSeek V4.1 Flash" for item in deepseek.models)
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    monkeypatch.setattr(
        ai_factory,
        "_run_security",
        lambda arguments: CompletedProcess(
            arguments, 0, stdout="secret-from-keychain\n", stderr=""
        ),
    )

    assert api_key_status("openrouter") == "Keychain"
    assert load_api_key("openrouter") == "secret-from-keychain"
    assert get_ai_provider("openrouter").name == "openrouter"


def test_windows_credential_manager_lookup(monkeypatch) -> None:
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    monkeypatch.setattr(ai_factory.sys, "platform", "win32")
    monkeypatch.setattr(ai_factory, "_load_windows_credential", lambda name: f"secret-{name}")

    assert api_key_status("openrouter") == "Windows Credential Manager"
    assert load_api_key("openrouter") == "secret-openrouter"


@pytest.mark.asyncio
async def test_provider_model_refresh_filters_non_text_models(monkeypatch) -> None:
    class FakeResponse:
        def raise_for_status(self) -> None:
            return None

        def json(self) -> dict:
            return {
                "data": [
                    {"id": "deepseek-flash", "name": "DeepSeek V4.1 Flash"},
                    {"id": "deepseek-v4-pro", "name": "DeepSeek V4 Pro"},
                    {"id": "deepseek-image", "name": "DeepSeek Image"},
                ]
            }

    class FakeClient:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def get(self, url, headers=None, params=None):
            assert url == "https://api.deepseek.com/v1/models"
            assert headers == {"Authorization": "Bearer secret"}
            return FakeResponse()

    monkeypatch.setattr(ai_factory, "load_api_key", lambda _name: "secret")
    monkeypatch.setattr(ai_factory.httpx, "AsyncClient", lambda **_kwargs: FakeClient())

    models = await ai_factory.fetch_provider_models("deepseek")

    assert {item.id for item in models} == {"deepseek-flash", "deepseek-v4-pro"}


def test_ai_note_exports_markdown_pdf_and_word(tmp_path: Path) -> None:
    context = LectureContext(
        course_code="SDDM6321",
        topic_name="Periodontics",
        lecture_number="1",
        title="Periodontal Maintenance",
        source="Echo360",
        source_recording_id="rec-1",
    )
    note = """---
course: SDDM6321
---

# Periodontal Maintenance

## Key points

- Review plaque control
    - **Monitor** periodontal stability
      - Record bleeding on probing

    | Finding | Interpretation |
    | --- | --- |
    | Stable pockets | Continue maintenance |
    | New bleeding | Reassess plaque control |
"""

    markdown = export_ai_note(note, tmp_path, context, "markdown")
    pdf = export_ai_note(note, tmp_path, context, "pdf")
    word = export_ai_note(note, tmp_path, context, "docx")
    vault_pdf = export_ai_note(
        note,
        tmp_path / "vault",
        context,
        "pdf",
        filename_stem=context.title,
    )

    assert markdown.read_text(encoding="utf-8") == note
    assert pdf.read_bytes().startswith(b"%PDF-1.4")
    assert b"Periodontal Maintenance" in pdf.read_bytes()
    assert word.read_bytes().startswith(b"PK")
    assert vault_pdf.name == "Periodontal Maintenance.pdf"
    document_object = Document(word)
    assert len(document_object.tables) == 1
    assert document_object.tables[0].cell(1, 0).text == "Stable pockets"
    assert any("Record bleeding on probing" in paragraph.text for paragraph in document_object.paragraphs)
    with zipfile.ZipFile(word) as archive:
        document = archive.read("word/document.xml").decode("utf-8")
    assert "Periodontal Maintenance" in document
    assert "course: SDDM6321" not in document


def test_echo360_syllabus_parser_uses_lessons_and_groups() -> None:
    payload = {
        "responses": [
            {
                "url": "https://echo.example/section/section-1/syllabus",
                "sample": {
                    "status": "ok",
                    "data": [
                        {
                            "groupInfo": {"name": "Tooth Conservation"},
                            "lessons": [
                                {
                                    "lesson": {
                                        "lesson": {
                                            "id": "lesson-1",
                                            "displayName": "Cracked Teeth",
                                            "timing": {"start": "2026-04-02T17:00:00"},
                                        },
                                        "medias": [
                                            {"id": "media-1", "title": "video.mp4"},
                                            {
                                                "id": "slides-1",
                                                "title": "Cracked Teeth.pdf",
                                                "mediaType": "Presentation",
                                                "isAvailable": True,
                                            },
                                        ],
                                    }
                                }
                            ],
                        }
                    ],
                },
            }
        ]
    }

    recordings = Echo360Client(load_settings()).parse_browser_payload(payload)

    assert len(recordings) == 1
    assert recordings[0].source_id == "lesson-1"
    assert recordings[0].title == "Cracked Teeth"
    assert recordings[0].topic_hint == "Tooth Conservation"
    assert recordings[0].url == "https://echo.example/lesson/lesson-1/classroom"
    assert recordings[0].presentations == [
        EchoPresentationAsset(media_id="slides-1", filename="Cracked Teeth.pdf")
    ]


@pytest.mark.asyncio
async def test_discovery_restarts_lecture_numbers_in_each_echo_folder(tmp_path: Path) -> None:
    db = Database(tmp_path / "app.sqlite3")
    settings = load_settings()
    db.upsert_course(Course(canvas_id="1", code="SDDM6321", name="Dental"))
    service = DiscoveryService(db, settings)
    service.canvas.list_modules = AsyncMock(return_value=[])
    service.canvas.echo360_navigation_links = AsyncMock(
        return_value=["https://echo.example/launch"]
    )
    service.echo.discover_course = AsyncMock(
        return_value=[
            EchoRecording(source_id="a-1", title="A One", url="https://echo/a-1", topic_hint="A"),
            EchoRecording(source_id="a-2", title="A Two", url="https://echo/a-2", topic_hint="A"),
            EchoRecording(
                source_id="b-1",
                title="B One",
                url="https://echo/b-1",
                topic_hint="B",
                presentations=[EchoPresentationAsset(media_id="slides-b1", filename="B One.pptx")],
            ),
        ]
    )

    await service.discover_course("SDDM6321", diagnostic_dir=tmp_path)

    course = db.get_course_by_code("SDDM6321")
    topics = {topic.name: topic for topic in db.list_topics(course.id or 0)}
    numbers = {
        topic_name: [
            lecture.lecture_number
            for lecture in db.list_lectures(course.id or 0)
            if lecture.topic_id == topic.id
        ]
        for topic_name, topic in topics.items()
    }
    assert numbers == {"A": ["1", "2"], "B": ["1"]}
    b_lecture = next(
        lecture for lecture in db.list_lectures(course.id or 0) if lecture.source_id == "b-1"
    )
    assert db.list_assets(b_lecture.id or 0)[0].remote_url == "echo360-presentation://slides-b1"


def test_echo360_asset_selection_prefers_highest_matching_tracks() -> None:
    tracks = [
        (1, 0, "https://content.example/s1q0.mp4"),
        (0, 1, "https://content.example/s0q1.mp4"),
        (1, 1, "https://content.example/s1q1.mp4"),
    ]

    assert _best_track_url(tracks, stream=1) == "https://content.example/s1q1.mp4"
    assert _best_track_url(tracks, stream=0) == "https://content.example/s0q1.mp4"
    assert (
        _matching_transcript_url(
            ["https://captions.example/intro.vtt", "https://captions.example/media-1.vtt"],
            ["media-1"],
        )
        == "https://captions.example/media-1.vtt"
    )


@pytest.mark.asyncio
async def test_canvas_finds_recorded_lectures_course_navigation(monkeypatch) -> None:
    settings = load_settings()
    settings.canvas.base_url = "https://canvas.example"
    client = CanvasClient(settings)
    browser_api = AsyncMock(
        return_value=[
            {
                "id": "context_external_tool_11653",
                "label": "Recorded Lectures",
                "type": "external",
                "html_url": "/courses/69082/external_tools/11653",
            }
        ]
    )
    monkeypatch.setattr(client, "_api_get_with_browser_session", browser_api)

    links = await client.echo360_navigation_links("69082")

    assert links == ["https://canvas.example/courses/69082/external_tools/11653"]


def test_database_idempotency_and_manifest(tmp_path: Path) -> None:
    db = Database(tmp_path / "app.sqlite3")
    course = db.upsert_course(
        Course(
            canvas_id="1",
            code="SDDM6321",
            name="Dental",
            group_name="Clinical Dentistry",
            is_active=True,
        )
    )
    course2 = db.upsert_course(Course(canvas_id="1", code="SDDM6321", name="Dental Updated"))
    assert course.id == course2.id
    assert course2.name == "Dental Updated"
    assert course2.group_name == "Clinical Dentistry"
    assert not course2.is_active
    assert db.update_course_groups({"1": "Dentistry"}) == 1
    topic = db.upsert_topic(
        Topic(course_id=course.id or 0, name="Endodontics", numbering_prefix="3")
    )
    lecture = db.upsert_lecture(
        Lecture(
            source_id="rec-1",
            course_id=course.id or 0,
            topic_id=topic.id,
            original_title="Intro",
            lecture_number="1",
        )
    )
    db.upsert_asset(
        Asset(lecture_id=lecture.id or 0, kind="video", status=DownloadStatus.UNAVAILABLE)
    )
    manifest_path = export_topic_manifest(db, topic.id or 0, tmp_path)
    assert manifest_path.exists()
    assert "Dental Updated" in manifest_path.read_text(encoding="utf-8")


def test_database_tracks_background_job_progress(tmp_path: Path) -> None:
    db = Database(tmp_path / "app.sqlite3")
    job = db.create_job(
        JobRecord(id="job-1", kind="discovery", course_code="SDDM6321", total_steps=6)
    )

    updated = db.update_job(
        job.id,
        status="running",
        current_step="Reading Echo360 folders and recordings",
        completed_steps=3,
    )

    assert updated.status == "running"
    assert updated.completed_steps == 3
    assert db.get_job(job.id) == updated


@pytest.mark.asyncio
async def test_pipeline_stage_skips_existing_transcript(tmp_path: Path) -> None:
    db = Database(tmp_path / "app.sqlite3")
    settings = load_settings()
    settings.paths.download_root = str(tmp_path)
    settings.paths.notebooklm_staging_dir = str(tmp_path / "notebook")
    course = db.upsert_course(Course(canvas_id="1", code="SDDM6321", name="Dental"))
    topic = db.upsert_topic(
        Topic(course_id=course.id or 0, name="Endodontics", numbering_prefix="3")
    )
    lecture = db.upsert_lecture(
        Lecture(
            source_id="rec-1",
            course_id=course.id or 0,
            topic_id=topic.id,
            original_title="Intro",
            lecture_number="1",
        )
    )
    topic_dir = tmp_path / "SDDM6321 Recorded Lectures" / "Endodontics"
    topic_dir.mkdir(parents=True)
    (topic_dir / "SDDM6321 Endodontics 1 Intro.txt").write_text("existing", encoding="utf-8")
    progress = []
    result = await PipelineService(db, settings).run_lecture(
        lecture.id or 0,
        PipelineOptions(download_video=False, download_transcript=True, notebooklm=True),
        progress=lambda message, completed, total: progress.append((message, completed, total)),
    )
    assert result["transcript"] == "skipped"
    assert Path(result["notebooklm"]).exists()
    assert progress[-1][1:] == (5, 5)
    assert any("Preparing transcript" in message for message, _, _ in progress)


@pytest.mark.asyncio
async def test_pipeline_downloads_slides_with_unified_name(tmp_path: Path) -> None:
    db = Database(tmp_path / "app.sqlite3")
    settings = load_settings()
    settings.paths.download_root = str(tmp_path)
    course = db.upsert_course(Course(canvas_id="1", code="SDDM6321", name="Dental"))
    topic = db.upsert_topic(Topic(course_id=course.id or 0, name="Periodontics"))
    lecture = db.upsert_lecture(
        Lecture(
            source_id="rec-1",
            course_id=course.id or 0,
            topic_id=topic.id,
            original_title="Periodontal Maintenance",
            lecture_number="1",
            recording_url="https://echo.example/lesson/rec-1/classroom",
        )
    )
    db.upsert_asset(
        Asset(
            lecture_id=lecture.id or 0,
            kind="slides",
            remote_url="echo360-presentation://slides-1",
            status=DownloadStatus.AVAILABLE,
        )
    )
    service = PipelineService(db, settings)
    service._resolve_echo_assets = AsyncMock(
        return_value=ResolvedEchoAssets(
            slides_url="https://echo.example/media/download/slides-1/slides.pptx",
            slides_extension="pptx",
        )
    )

    async def download(_url, destination, **_kwargs):
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(b"slides")
        return "complete"

    service.downloads.download = AsyncMock(side_effect=download)
    result = await service.run_lecture(
        lecture.id or 0,
        PipelineOptions(
            download_video=False,
            download_slides=True,
            download_transcript=False,
            notebooklm=False,
        ),
    )

    expected = (
        tmp_path
        / "SDDM6321 Recorded Lectures"
        / "Periodontics"
        / "SDDM6321 Periodontics 1 Periodontal Maintenance.pptx"
    )
    assert result["slides"] == "complete"
    assert expected.read_bytes() == b"slides"


@pytest.mark.asyncio
async def test_pipeline_progress_identifies_each_lecture(tmp_path: Path) -> None:
    db = Database(tmp_path / "app.sqlite3")
    settings = load_settings()
    settings.paths.download_root = str(tmp_path)
    course = db.upsert_course(Course(canvas_id="1", code="SDDM6321", name="Dental"))
    topic = db.upsert_topic(Topic(course_id=course.id or 0, name="Radiology"))
    first = db.upsert_lecture(
        Lecture(
            source_id="rec-1", course_id=course.id or 0, topic_id=topic.id, original_title="TMJ"
        )
    )
    second = db.upsert_lecture(
        Lecture(
            source_id="rec-2", course_id=course.id or 0, topic_id=topic.id, original_title="Trauma"
        )
    )
    progress = []

    await PipelineService(db, settings).run_many(
        [first.id or 0, second.id or 0],
        PipelineOptions(
            download_video=False,
            download_transcript=False,
            notebooklm=False,
        ),
        progress=lambda message, completed, total: progress.append((message, completed, total)),
    )

    messages = [message for message, _, _ in progress]
    assert any(message == "Lecture 1 of 2: Preparing - TMJ" for message in messages)
    assert any(
        message == "Lecture 2 of 2: Updating lecture manifest - Trauma" for message in messages
    )
    assert progress[-1][1:] == (4, 4)


@pytest.mark.asyncio
async def test_pipeline_continues_after_one_lecture_fails(tmp_path: Path) -> None:
    db = Database(tmp_path / "app.sqlite3")
    settings = load_settings()
    course = db.upsert_course(Course(canvas_id="1", code="DENT1000", name="Dental"))
    topic = db.upsert_topic(Topic(course_id=course.id or 0, name="Oral Medicine"))
    first = db.upsert_lecture(
        Lecture(
            source_id="rec-1",
            course_id=course.id or 0,
            topic_id=topic.id,
            original_title="First",
        )
    )
    second = db.upsert_lecture(
        Lecture(
            source_id="rec-2",
            course_id=course.id or 0,
            topic_id=topic.id,
            original_title="Second",
        )
    )
    service = PipelineService(db, settings)
    service.run_lecture = AsyncMock(
        side_effect=[RuntimeError("first failed"), {"lecture": "Second", "transcript": "complete"}]
    )
    progress = []

    results = await service.run_many(
        [first.id or 0, second.id or 0],
        PipelineOptions(
            download_video=False,
            download_transcript=False,
            notebooklm=False,
        ),
        progress=lambda message, completed, total: progress.append((message, completed, total)),
    )

    assert results[0] == {"lecture": "First", "error": "first failed"}
    assert results[1]["lecture"] == "Second"
    assert service.run_lecture.await_count == 2
    assert db.list_recent_processing(course.id or 0)[0]["stage"] == "pipeline"
    assert progress[0][1:] == (2, 4)
