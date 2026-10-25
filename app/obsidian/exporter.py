from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import UTC, datetime
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt
from markdown_it.tree import SyntaxTreeNode
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    LongTable,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    TableStyle,
)

from app.models import LectureContext
from app.utils import lecture_stem, safe_filename, sha256_text

AI_NOTE_FORMATS = {"markdown": ".md", "pdf": ".pdf", "docx": ".docx"}


def obsidian_note_markdown(
    context: LectureContext,
    body: str,
    *,
    parent_note: str | None,
    ai_provider: str,
    ai_model: str,
    prompt_version: str,
    transcript: str,
) -> str:
    up = f'"[[{parent_note}]]"' if parent_note else '""'
    return f"""---
up: {up}
course: {context.course_code}
lecture: "{context.lecture_number}"
source: {context.source}
ai_generated: true
ai_provider: "{ai_provider}"
ai_model: "{ai_model}"
prompt_version: "{prompt_version}"
transcript_hash: "{sha256_text(transcript)}"
generated_at: "{datetime.now(UTC).isoformat()}"
---

# {context.title}

{body.strip()}
"""


def save_obsidian_note(
    markdown: str, destination_dir: Path, context: LectureContext, overwrite: bool = False
) -> Path:
    destination_dir.mkdir(parents=True, exist_ok=True)
    path = destination_dir / f"{safe_filename(context.title)}.md"
    if path.exists() and not overwrite:
        return path
    tmp = path.with_suffix(path.suffix + ".part")
    tmp.write_text(markdown, encoding="utf-8")
    tmp.replace(path)
    return path


def export_ai_note(
    markdown: str,
    destination_dir: Path,
    context: LectureContext,
    note_format: str,
    overwrite: bool = False,
    filename_stem: str | None = None,
) -> Path:
    suffix = AI_NOTE_FORMATS.get(note_format)
    if not suffix:
        raise ValueError(f"Unsupported AI note format: {note_format}")
    destination_dir.mkdir(parents=True, exist_ok=True)
    filename = filename_stem or lecture_stem(
        context.course_code, context.topic_name, context.lecture_number, context.title
    )
    path = destination_dir / f"{filename}{suffix}"
    if path.exists() and not overwrite:
        return path
    if note_format == "markdown":
        payload = markdown.encode("utf-8")
    elif note_format == "docx":
        payload = _docx_bytes(markdown, context)
    else:
        payload = _pdf_bytes(markdown, context)
    tmp = path.with_suffix(path.suffix + ".part")
    tmp.write_bytes(payload)
    tmp.replace(path)
    return path


def _note_body(markdown: str) -> str:
    if not markdown.startswith("---\n"):
        return markdown.strip()
    end = markdown.find("\n---\n", 4)
    return markdown[end + 5 :].strip() if end >= 0 else markdown.strip()


@dataclass(frozen=True)
class InlineSpan:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    href: str | None = None


def _markdown_tree(markdown: str) -> SyntaxTreeNode:
    parser = MarkdownIt("commonmark", {"html": False}).enable("table")
    return SyntaxTreeNode(parser.parse(_note_body(markdown)))


def _inline_spans(
    node: SyntaxTreeNode,
    *,
    bold: bool = False,
    italic: bool = False,
    code: bool = False,
    href: str | None = None,
) -> list[InlineSpan]:
    node_type = node.type
    bold = bold or node_type == "strong"
    italic = italic or node_type == "em"
    code = code or node_type == "code_inline"
    href = str(node.attrs.get("href")) if node_type == "link" else href
    if node_type in {"text", "code_inline"}:
        return [InlineSpan(node.content, bold, italic, code, href)]
    if node_type in {"softbreak", "hardbreak"}:
        return [InlineSpan("\n", bold, italic, code, href)]
    if node_type == "image":
        return [InlineSpan(f"[{node.content or 'Image'}]", bold, italic, code, href)]
    spans: list[InlineSpan] = []
    for child in node.children:
        spans.extend(_inline_spans(child, bold=bold, italic=italic, code=code, href=href))
    return spans


def _first_inline(node: SyntaxTreeNode) -> list[InlineSpan]:
    inline = next((child for child in node.children if child.type == "inline"), None)
    return _inline_spans(inline) if inline else []


def _add_docx_spans(paragraph, spans: list[InlineSpan]) -> None:
    for span in spans:
        text = f"{span.text} ({span.href})" if span.href else span.text
        run = paragraph.add_run(text)
        run.bold = span.bold
        run.italic = span.italic
        if span.code:
            run.font.name = "Menlo"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(36, 56, 70)


def _shade_docx(element, fill: str) -> None:
    properties = (
        element._tc.get_or_add_tcPr() if hasattr(element, "_tc") else element._p.get_or_add_pPr()
    )
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _repeat_docx_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _render_docx_list(document: Document, node: SyntaxTreeNode, level: int = 0) -> None:
    style = "List Number" if node.type == "ordered_list" else "List Bullet"
    for item in (child for child in node.children if child.type == "list_item"):
        for child in item.children:
            if child.type == "paragraph":
                paragraph = document.add_paragraph(style=style)
                paragraph.paragraph_format.left_indent = Inches(0.2 * level)
                _add_docx_spans(paragraph, _first_inline(child))
            elif child.type in {"bullet_list", "ordered_list"}:
                _render_docx_list(document, child, level + 1)
            else:
                _render_docx_node(document, child)


def _render_docx_table(document: Document, node: SyntaxTreeNode) -> None:
    rows = [row for section in node.children for row in section.children if row.type == "tr"]
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=max(len(row.children) for row in rows))
    table.style = "Table Grid"
    table.autofit = True
    for row_index, source_row in enumerate(rows):
        for column_index, source_cell in enumerate(source_row.children):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _add_docx_spans(cell.paragraphs[0], _first_inline(source_cell))
            if row_index == 0:
                _shade_docx(cell, "DCE6F1")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
        if row_index == 0:
            _repeat_docx_header(table.rows[0])


def _render_docx_node(document: Document, node: SyntaxTreeNode) -> None:
    if node.type == "heading":
        paragraph = document.add_heading(level=min(3, max(1, int(node.tag[1:]))))
        _add_docx_spans(paragraph, _first_inline(node))
    elif node.type == "paragraph":
        paragraph = document.add_paragraph()
        _add_docx_spans(paragraph, _first_inline(node))
    elif node.type in {"bullet_list", "ordered_list"}:
        _render_docx_list(document, node)
    elif node.type == "table":
        _render_docx_table(document, node)
        document.add_paragraph()
    elif node.type == "blockquote":
        for child in node.children:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.right_indent = Inches(0.2)
            _add_docx_spans(paragraph, _first_inline(child))
            _shade_docx(paragraph, "EEF3F7")
    elif node.type in {"fence", "code_block"}:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.2)
        run = paragraph.add_run(node.content.rstrip())
        run.font.name = "Menlo"
        run.font.size = Pt(8.5)
        _shade_docx(paragraph, "F1F4F6")
    else:
        for child in node.children:
            _render_docx_node(document, child)


def _add_docx_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _docx_bytes(markdown: str, context: LectureContext) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    for level, size, color in ((1, 20, "113F67"), (2, 15, "0B6E4F"), (3, 12, "365469")):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if level > 1 else 0)
        style.paragraph_format.space_after = Pt(6)
    header = section.header.paragraphs[0]
    header.text = f"Canvas Lecture AI  |  {context.course_code}  |  {context.topic_name}"
    header.style = document.styles["Caption"]
    _add_docx_page_number(section.footer.paragraphs[0])
    document.core_properties.title = context.title
    document.core_properties.subject = f"{context.course_code} - {context.topic_name}"
    document.core_properties.author = "Canvas Lecture AI"
    for node in _markdown_tree(markdown).children:
        _render_docx_node(document, node)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _register_pdf_fonts() -> tuple[str, str, str, str]:
    try:
        import reportlab

        font_dir = Path(reportlab.__file__).parent / "fonts"
        font_files = (
            ("CanvasSans", "Vera.ttf"),
            ("CanvasSans-Bold", "VeraBd.ttf"),
            ("CanvasSans-Italic", "VeraIt.ttf"),
        )
        for name, filename in font_files:
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(font_dir / filename)))
        return "CanvasSans", "CanvasSans-Bold", "CanvasSans-Italic", "Courier"
    except (OSError, ValueError):
        return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Courier"


def _pdf_inline(spans: list[InlineSpan], fonts: tuple[str, str, str, str]) -> str:
    output: list[str] = []
    for span in spans:
        value = escape(span.text).replace("\n", "<br/>")
        font = fonts[3] if span.code else fonts[0]
        font = fonts[1] if span.bold else fonts[2] if span.italic else font
        value = f'<font name="{font}">{value}</font>'
        if span.href:
            value = f'<link href="{escape(span.href)}" color="#0B6E4F">{value}</link>'
        output.append(value)
    return "".join(output)


def _pdf_styles(fonts: tuple[str, str, str, str]) -> dict[str, ParagraphStyle]:
    sample = getSampleStyleSheet()
    body = ParagraphStyle(
        "CanvasBody",
        parent=sample["BodyText"],
        fontName=fonts[0],
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#1F2933"),
        spaceAfter=5,
    )
    return {
        "body": body,
        "h1": ParagraphStyle(
            "CanvasH1",
            parent=sample["Heading1"],
            fontName=fonts[1],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor("#113F67"),
            spaceAfter=10,
        ),
        "h2": ParagraphStyle(
            "CanvasH2",
            parent=sample["Heading2"],
            fontName=fonts[1],
            fontSize=14,
            leading=18,
            textColor=colors.HexColor("#0B6E4F"),
            spaceBefore=10,
            spaceAfter=5,
            keepWithNext=True,
        ),
        "h3": ParagraphStyle(
            "CanvasH3",
            parent=sample["Heading3"],
            fontName=fonts[1],
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#365469"),
            spaceBefore=8,
            spaceAfter=4,
            keepWithNext=True,
        ),
        "quote": ParagraphStyle(
            "CanvasQuote",
            parent=body,
            leftIndent=10,
            rightIndent=8,
            borderColor=colors.HexColor("#7EA4BF"),
            borderWidth=1,
            borderPadding=6,
            backColor=colors.HexColor("#EEF3F7"),
        ),
        "code": ParagraphStyle(
            "CanvasCode",
            parent=sample["Code"],
            fontName=fonts[3],
            fontSize=7.8,
            leading=10,
            leftIndent=8,
            borderPadding=6,
            backColor=colors.HexColor("#F1F4F6"),
        ),
        "table": ParagraphStyle(
            "CanvasTable", parent=body, fontSize=7.5, leading=9.5, spaceAfter=0
        ),
        "table_header": ParagraphStyle(
            "CanvasTableHeader",
            parent=body,
            fontName=fonts[1],
            fontSize=7.5,
            leading=9.5,
            textColor=colors.HexColor("#113F67"),
            spaceAfter=0,
        ),
    }


def _pdf_list(node: SyntaxTreeNode, styles, fonts, level: int = 0):
    items = []
    for item in (child for child in node.children if child.type == "list_item"):
        flows = []
        for child in item.children:
            if child.type == "paragraph":
                flows.append(Paragraph(_pdf_inline(_first_inline(child), fonts), styles["body"]))
            elif child.type in {"bullet_list", "ordered_list"}:
                flows.append(_pdf_list(child, styles, fonts, level + 1))
            else:
                flows.extend(_pdf_node(child, styles, fonts))
        items.append(ListItem(flows or [Spacer(1, 1)], leftIndent=8 + level * 4))
    return ListFlowable(
        items,
        bulletType="1" if node.type == "ordered_list" else "bullet",
        leftIndent=18 + level * 6,
        bulletFontName=fonts[0],
        bulletFontSize=8,
        spaceAfter=5,
    )


def _pdf_table(node: SyntaxTreeNode, styles, fonts):
    rows = [row for section in node.children for row in section.children if row.type == "tr"]
    data = []
    for row_index, row in enumerate(rows):
        style = styles["table_header"] if row_index == 0 else styles["table"]
        data.append(
            [Paragraph(_pdf_inline(_first_inline(cell), fonts), style) for cell in row.children]
        )
    if not data:
        return Spacer(1, 1)
    columns = max(len(row) for row in data)
    usable_width = A4[0] - 32 * mm
    table = LongTable(
        data, colWidths=[usable_width / columns] * columns, repeatRows=1, splitByRow=1
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B8C5D1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def _pdf_node(node: SyntaxTreeNode, styles, fonts) -> list:
    if node.type == "heading":
        level = min(3, max(1, int(node.tag[1:])))
        return [Paragraph(_pdf_inline(_first_inline(node), fonts), styles[f"h{level}"])]
    if node.type == "paragraph":
        return [Paragraph(_pdf_inline(_first_inline(node), fonts), styles["body"])]
    if node.type in {"bullet_list", "ordered_list"}:
        return [_pdf_list(node, styles, fonts)]
    if node.type == "table":
        return [_pdf_table(node, styles, fonts), Spacer(1, 5)]
    if node.type == "blockquote":
        spans = [span for child in node.children for span in _first_inline(child)]
        return [Paragraph(_pdf_inline(spans, fonts), styles["quote"])]
    if node.type in {"fence", "code_block"}:
        return [Preformatted(node.content.rstrip(), styles["code"])]
    flows = []
    for child in node.children:
        flows.extend(_pdf_node(child, styles, fonts))
    return flows


def _pdf_bytes(markdown: str, context: LectureContext) -> bytes:
    output = io.BytesIO()
    fonts = _register_pdf_fonts()
    styles = _pdf_styles(fonts)
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=18 * mm,
        bottomMargin=17 * mm,
        title=context.title,
        author="Canvas Lecture AI",
        subject=f"{context.course_code} - {context.topic_name}",
    )

    def page(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont(fonts[0], 7.5)
        canvas.setFillColor(colors.HexColor("#526574"))
        canvas.drawString(16 * mm, A4[1] - 10 * mm, f"Canvas Lecture AI  |  {context.course_code}")
        canvas.drawRightString(A4[0] - 16 * mm, 9 * mm, f"Page {doc.page}")
        canvas.restoreState()

    story = [
        flow
        for node in _markdown_tree(markdown).children
        for flow in _pdf_node(node, styles, fonts)
    ]
    if not story:
        story = [Paragraph("No note content was generated.", styles["body"])]
    document.build(story, onFirstPage=page, onLaterPages=page)
    return output.getvalue()
